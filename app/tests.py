import re
from pathlib import Path

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from . import catalog
from .models import BabyCotModel, Enquiry, Job

PAGES = [
    "Home", "products", "BabyCot", "ClassicRanges", "SkinRanges", "VeneerRanges", "LaminateRanges",
    "LightOpeningRanges", "Furnitureforkids", "Tvunit", "Wardrobe", "Corporate", "seasoningchamber",
    "careers", "contact", "contact_thanks",
]


class PageTests(TestCase):
    def test_every_page_loads(self):
        for name in PAGES:
            r = self.client.get(reverse(name))
            self.assertEqual(r.status_code, 200, name)

    def test_menu_is_unchanged_on_every_page(self):
        for name in PAGES:
            html = self.client.get(reverse(name)).content.decode()
            for label in ["Home", "About", "Our Products", "Our Client", "Contact", "Residence", "Our facility", "Careers"]:
                self.assertIn(label, html, f"{label} missing on {name}")
            self.assertNotIn("WhatsApp", html)

    def test_footer_links_are_correct(self):
        html = self.client.get(reverse("Home")).content.decode()
        self.assertIn('href="tel:+918122507384"', html)
        self.assertIn("mailto:neminathwoodindustry.pvt.ltd@gmail.com", html)
        self.assertNotIn("info@company.com", html)
        self.assertNotIn("2027", html)

    def test_baby_cot_models_seeded(self):
        self.assertEqual(BabyCotModel.objects.count(), 4)
        self.assertContains(self.client.get(reverse("BabyCot")), "NWI-BC-1A-04")

    def test_404(self):
        self.assertEqual(self.client.get("/does-not-exist/").status_code, 404)

    def test_sitemap_and_robots(self):
        self.assertEqual(self.client.get("/sitemap.xml").status_code, 200)
        self.assertEqual(self.client.get("/robots.txt").status_code, 200)


class EnquiryTests(TestCase):
    data = {"name": "Ravi", "phone": "+91 98765 43210", "email": "", "enquiry_type": "project",
            "message": "Need 40 doors", "source": "Corporate solutions", "website": ""}

    def test_enquiry_is_saved(self):
        r = self.client.post(reverse("contact"), self.data)
        self.assertRedirects(r, reverse("contact_thanks"))
        e = Enquiry.objects.get()
        self.assertEqual((e.name, e.source), ("Ravi", "Corporate solutions"))

    def test_honeypot_blocks_bots(self):
        r = self.client.post(reverse("contact"), {**self.data, "website": "http://spam"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(Enquiry.objects.count(), 0)

    def test_bad_phone_rejected(self):
        r = self.client.post(reverse("contact"), {**self.data, "phone": "abc"})
        self.assertEqual(r.status_code, 400)
        self.assertContains(r, "valid phone", status_code=400)

    def test_rate_limit(self):
        from django.core.cache import cache
        cache.clear()
        for _ in range(5):
            self.client.post(reverse("contact"), self.data)
        self.assertEqual(self.client.post(reverse("contact"), self.data).status_code, 429)
        cache.clear()

    def test_ref_prefills_form(self):
        self.assertContains(self.client.get(reverse("contact") + "?ref=Wardrobe"), "enquire about: Wardrobe")


class CareerTests(TestCase):
    def setUp(self):
        self.job = Job.objects.create(title="CNC Operator", department="Production", location="Chennai",
                                      experience_required="2 years", job_type="Full Time", description="Run CNC machines.")

    def test_list_detail_apply_pages(self):
        self.assertContains(self.client.get(reverse("careers")), "CNC Operator")
        self.assertContains(self.client.get(reverse("job_detail", args=[self.job.id])), "Run CNC machines.")
        self.assertContains(self.client.get(reverse("job_apply", args=[self.job.id])), "Personal Details")

    def test_apply_saves(self):
        import tempfile
        from django.test import override_settings
        with tempfile.TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp):
            r = self.client.post(reverse("job_apply", args=[self.job.id]), {
                "first_name": "A", "last_name": "B", "email": "a@b.com", "contact": "9876543210",
                "year_of_graduation": 2020, "gender": "Male", "experience_years": "2.0",
                "current_employer": "", "current_ctc": "3.5", "expected_ctc": "4.5",
                "notice_period": "30 Days", "current_location": "Chennai",
                "resume": SimpleUploadedFile("cv.pdf", b"%PDF-1.4 test", content_type="application/pdf"),
            })
            self.assertRedirects(r, reverse("careers"))
            self.assertEqual(self.job.applications.count(), 1)

    def test_inactive_job_hidden(self):
        self.job.is_active = False
        self.job.save()
        self.assertEqual(self.client.get(reverse("job_detail", args=[self.job.id])).status_code, 404)


class StaticReferenceTests(TestCase):
    def test_every_image_exists(self):
        static = Path(settings.BASE_DIR) / "static"
        paths = set()
        for s in catalog.HERO_SLIDES:
            paths.add(s)
        for _, p, _ in catalog.SHOP_TILES:
            paths.add("images/" + p)
        paths.update(c[2] for c in catalog.COLLECTIONS)
        paths.update(f[2] for f in catalog.PRODUCT_FEATURES)
        paths.update(g[0] for g in catalog.PRODUCT_GALLERY)
        paths.add(catalog.BABYCOT_INTRO["image"])
        for r in catalog.RANGES.values():
            paths.update([r["intro_img"], r["about_img"]] + [g[0] for g in r["gallery"]])
        for s in catalog.SPACES.values():
            paths.update([s["a_img"], s["b_img"]] + [g[0] for g in s["gallery"]])
        paths.add(catalog.SEASONING["hero_img"])
        paths.update(st[2] for st in catalog.SEASONING["steps"] if st[2])
        paths.update("images/clientlogos/" + f for _, f in catalog.CLIENTS)
        paths.update(m.image for m in BabyCotModel.objects.all())
        paths.update(["images/NeminathLogo.jpeg", "images/logo-mark.jpg", "images/logo-full-crop.jpg",
                      "images/NWIPLLOGO.png", "images/background2.jpg", "images/featured-compressed-46.jpg",
                      "downloads/NeminathProfileBrochure.pdf"])
        missing = sorted(p for p in paths if not (static / p).exists())
        self.assertEqual(missing, [])
        self.assertEqual(len(catalog.CLIENTS), 22)


class ExcelExportTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.admin = get_user_model().objects.create_superuser("boss", "b@x.com", "pw12345!")
        Enquiry.objects.create(name="Ravi", phone="+91 98765 43210", enquiry_type="project",
                               message="=SUM(1+1)", source="Corporate solutions")
        Enquiry.objects.create(name="Meena", phone="9000000000", enquiry_type="home", is_handled=True)

    def _load(self, response):
        import io
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(response.content))

    def test_button_downloads_excel(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("admin:app_enquiry_export"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn(".xlsx", r["Content-Disposition"])
        wb = self._load(r)
        ws = wb["Enquiries"]
        self.assertEqual([c.value for c in ws[1]][:4], ["Date", "Time", "Name", "Phone"])
        self.assertEqual(ws.max_row, 3)
        names = {ws.cell(row=i, column=3).value for i in (2, 3)}
        self.assertEqual(names, {"Ravi", "Meena"})
        self.assertEqual(wb["Per day"].cell(row=2, column=2).value, 2)

    def test_formula_text_is_not_executed(self):
        import io
        import zipfile
        self.client.force_login(self.admin)
        r = self.client.get(reverse("admin:app_enquiry_export"))
        ws = self._load(r)["Enquiries"]
        self.assertIn("=SUM(1+1)", [ws.cell(row=i, column=8).value for i in (2, 3)])
        sheet_xml = zipfile.ZipFile(io.BytesIO(r.content)).read("xl/worksheets/sheet1.xml").decode()
        self.assertNotIn("<f>", sheet_xml)  # no real formula cells

    def test_filter_applies(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("admin:app_enquiry_export") + "?enquiry_type__exact=project")
        self.assertEqual(self._load(r)["Enquiries"].max_row, 2)

    def test_changelist_shows_button_and_action(self):
        self.client.force_login(self.admin)
        html = self.client.get(reverse("admin:app_enquiry_changelist")).content.decode()
        self.assertIn("Download Excel", html)
        self.assertIn("Download selected enquiries as Excel", html)

    def test_selected_action(self):
        self.client.force_login(self.admin)
        ids = list(Enquiry.objects.values_list("pk", flat=True))
        r = self.client.post(reverse("admin:app_enquiry_changelist"),
                             {"action": "download_excel", "_selected_action": ids[:1]})
        self.assertEqual(self._load(r)["Enquiries"].max_row, 2)

    def test_public_cannot_download(self):
        r = self.client.get(reverse("admin:app_enquiry_export"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("login", r["Location"])
