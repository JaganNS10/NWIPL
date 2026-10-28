"""Excel export of enquiries (date-wise)."""
from collections import Counter

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Date", "Time", "Name", "Phone", "Email", "Type", "Enquired about", "Message", "Email sent", "Handled"]
WIDTHS = [14, 10, 24, 18, 30, 18, 28, 60, 11, 10]


def _put(ws, row, col, value):
    """Write a value; text is always stored as text so a message like '=SUM(1)' can never run as a formula."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def enquiries_excel_response(queryset):
    rows = list(queryset.order_by("-created_at"))
    wb = Workbook()

    ws = wb.active
    ws.title = "Enquiries"
    head_fill = PatternFill("solid", fgColor="3A2A14")
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="E3C37F")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]

    per_day = Counter()
    for r, e in enumerate(rows, start=2):
        local = timezone.localtime(e.created_at)
        per_day[local.date()] += 1
        d = ws.cell(row=r, column=1, value=local.date())
        d.number_format = "DD-MMM-YYYY"
        _put(ws, r, 2, local.strftime("%I:%M %p"))
        _put(ws, r, 3, e.name)
        _put(ws, r, 4, e.phone)
        _put(ws, r, 5, e.email)
        _put(ws, r, 6, e.get_enquiry_type_display())
        _put(ws, r, 7, e.source)
        m = _put(ws, r, 8, e.message)
        m.alignment = Alignment(wrap_text=True, vertical="top")
        _put(ws, r, 9, "Yes" if e.email_sent else "No")
        _put(ws, r, 10, "Yes" if e.is_handled else "No")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(len(rows) + 1, 2)}"

    summary = wb.create_sheet("Per day")
    for c, h in enumerate(["Date", "Number of enquiries"], start=1):
        cell = summary.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="E3C37F")
        cell.fill = head_fill
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 22
    for r, (day, n) in enumerate(sorted(per_day.items(), reverse=True), start=2):
        cell = summary.cell(row=r, column=1, value=day)
        cell.number_format = "DD-MMM-YYYY"
        summary.cell(row=r, column=2, value=n)
    summary.cell(row=len(per_day) + 3, column=1, value="Total").font = Font(bold=True)
    summary.cell(row=len(per_day) + 3, column=2, value=len(rows)).font = Font(bold=True)

    filename = f"NWI-enquiries-{timezone.localdate():%Y-%m-%d}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
